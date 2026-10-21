"""Exercise the analysis/PDF boundary with explicitly synthetic inputs."""
import argparse
import json
from pathlib import Path
import subprocess
import sys


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    directory = args.output.resolve()
    directory.mkdir(parents=True, exist_ok=True)
    import numpy as np
    import pandas as pd
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from rdkit import Chem
    from rdkit.Chem import AllChem
    import fitz
    from pypdf import PdfReader
    import pdfplumber
    from docx import Document

    data = pd.DataFrame({'x': [0, 1, 2, 3], 'y': np.square([0, 1, 2, 3])})
    data.to_csv(directory / 'synthetic.csv', index=False)
    figure, axes = plt.subplots()
    axes.plot(data.x, data.y)
    axes.set(title='Synthetic runtime verification', xlabel='x', ylabel='x squared')
    figure.savefig(directory / 'synthetic.png')
    plt.close(figure)
    mol = Chem.AddHs(Chem.MolFromSmiles('CCO'))
    assert AllChem.EmbedMolecule(mol, randomSeed=42) == 0
    (directory / 'ethanol.xyz').write_text(Chem.MolToXYZBlock(mol))
    pdf = fitz.open()
    page = pdf.new_page()
    title = 'EM-Agent synthetic runtime verification'
    page.insert_text((72, 72), title)
    pdf.save(directory / 'synthetic.pdf')
    pdf.close()
    assert title in PdfReader(directory / 'synthetic.pdf').pages[0].extract_text()
    with pdfplumber.open(directory / 'synthetic.pdf') as document:
        assert title in document.pages[0].extract_text()
    subprocess.run(['pdftotext', str(directory / 'synthetic.pdf'), str(directory / 'synthetic.txt')], check=True)
    assert title in (directory / 'synthetic.txt').read_text()
    subprocess.run(['pdftoppm', '-f', '1', '-singlefile', '-scale-to', '600', '-png',
                    str(directory / 'synthetic.pdf'), str(directory / 'pdf-page')], check=True)
    assert (directory / 'pdf-page.png').read_bytes().startswith(b'\x89PNG')
    document = Document()
    document.add_paragraph(title)
    document.save(directory / 'synthetic.docx')
    assert Document(directory / 'synthetic.docx').paragraphs[0].text == title
    data.to_excel(directory / 'synthetic.xlsx', index=False)
    assert pd.read_excel(directory / 'synthetic.xlsx').equals(data)
    result = {'verified': True, 'synthetic': True, 'python': sys.executable,
              'version': sys.version.split()[0], 'atoms': mol.GetNumAtoms(),
              'files': sorted(path.name for path in directory.iterdir())}
    (directory / 'verification.json').write_text(json.dumps(result, indent=2) + '\n')
    print(json.dumps(result))


if __name__ == '__main__':
    main()
